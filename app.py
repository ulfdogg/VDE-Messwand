"""
VDE Messwand - Hauptanwendung
"""
from flask import Flask, render_template, request, jsonify, Response, send_file
from jinja2 import FileSystemLoader
import os
import subprocess
import io
import csv
import time
from datetime import datetime

# Import eigener Module
from config import *
from managers.database import *
from hardware.relay_controller import RelayController
from managers.exam_utils import *
from managers.group_manager import *
from managers.settings_manager import *
from managers.stromkreis_manager import *
from managers.relais_manager import *
from managers.training_manager import *
from managers.relais_templates import *
from managers.relais_excel import *
from managers.network_manager import (
    is_hotspot_active, toggle_hotspot, get_wifi_networks,
    connect_to_wifi, get_current_connection, get_network_info,
    get_ethernet_info
)
from hardware.gpio_monitor import init_gpio_monitor, get_gpio_status, cleanup_gpio, update_gpio_shutdown_timeout

# Flask App initialisieren
app = Flask(__name__)
app.secret_key = SECRET_KEY
app.jinja_loader = FileSystemLoader('templates', encoding='utf-8')

# Logging-Filter für GPIO-Status API
import logging
class NoGPIOStatusFilter(logging.Filter):
    def filter(self, record):
        # Filtere /api/gpio/status Requests heraus
        return '/api/gpio/status' not in record.getMessage()

# Füge Filter zu Werkzeug-Logger hinzu
werkzeug_logger = logging.getLogger('werkzeug')
werkzeug_logger.addFilter(NoGPIOStatusFilter())

# Globale Instanzen
relay_controller = RelayController()
exam_active = False
exam_client_ip = None  # IP des Clients, der die Prüfung gestartet hat
exam_number_current = None
exam_start_time = None  # Unix-Timestamp (time.time())
# Sicherheitsrelais – immer AUS beim Start, Einschalten nur per Admin-Code
# Zustand wird im relay_controller.safety_relay_state gespiegelt


def get_client_ip():
    """Gibt die IP-Adresse des aktuellen Clients zurück (auch hinter Proxies)"""
    return request.headers.get('X-Forwarded-For', request.remote_addr).split(',')[0].strip()


def exam_lock_response():
    """Standardantwort wenn eine Prüfung aktiv ist und die Aktion blockiert wird"""
    return jsonify({
        'success': False,
        'error': 'Prüfung aktiv',
        'message': 'Diese Aktion ist während einer laufenden Prüfung gesperrt.'
    }), 423

# GPIO-Monitor initialisieren (Standard: GPIO 17 und 27)
# Kann in config.py angepasst werden
gpio_monitor_instance = None


# ==================== HAUPT-ROUTEN ====================

@app.route('/')
def index():
    """Startseite"""
    from managers.settings_manager import get_wallbox_enabled, get_wallbox_installed
    wallbox_installed = get_wallbox_installed()
    wallbox_enabled = get_wallbox_enabled()
    return render_template('index.html',
                         wallbox_installed=wallbox_installed,
                         wallbox_enabled=wallbox_enabled)


@app.route('/api/wallbox/toggle', methods=['POST'])
def api_toggle_wallbox():
    """API: Wallbox-Stromkreis ein/ausschalten"""
    from managers.settings_manager import set_wallbox_enabled

    data = request.json
    enabled = data.get('enabled', True)

    success, message = set_wallbox_enabled(enabled)

    return jsonify({
        'success': success,
        'message': message,
        'wallbox_enabled': enabled
    })


@app.route('/api/wallbox/installed', methods=['POST'])
def api_set_wallbox_installed():
    """API: Wallbox als vorhanden/nicht vorhanden markieren"""
    from managers.settings_manager import set_wallbox_installed

    data = request.json
    installed = data.get('installed', False)

    success, message = set_wallbox_installed(installed)

    return jsonify({
        'success': success,
        'message': message,
        'wallbox_installed': installed
    })


# ==================== PRÜFUNGSMODUS ====================

@app.route('/exam_mode')
def exam_mode():
    """Prüfungsmodus-Seite"""
    exam_settings = get_exam_settings()
    if exam_active and exam_number_current:
        number = exam_number_current
    else:
        number = generate_exam_number()
    return render_template('exam_mode.html',
                           exam_number=number,
                           exam_duration_minutes=exam_settings['exam_duration_minutes'])


@app.route('/start_exam', methods=['POST'])
def start_exam():
    """Startet eine neue Prüfung mit zufälligen Fehlern"""
    global exam_active, exam_client_ip, exam_number_current, exam_start_time
    import time as _time

    exam_number = request.json.get('exam_number')
    selected_relays = select_random_relays()

    # Relais aktivieren (Gruppen werden automatisch zusammen geschaltet)
    for relay_id in selected_relays:
        relay_controller.set_relay(relay_id, True)

    # In Datenbank speichern (wird automatisch normalisiert)
    save_examination(exam_number, selected_relays)

    exam_active = True
    exam_client_ip = get_client_ip()
    exam_number_current = exam_number
    exam_start_time = _time.time()
    print(f"🔒 Prüfung gestartet von IP: {exam_client_ip}")
    return jsonify({
        'success': True,
        'selected_errors': selected_relays,
        'exam_number': exam_number
    })


@app.route('/finish_exam', methods=['POST'])
def finish_exam():
    """Beendet eine Prüfung"""
    global exam_active, exam_client_ip, exam_number_current, exam_start_time

    # Nur die IP, die die Prüfung gestartet hat, darf sie beenden
    if exam_active and exam_client_ip and get_client_ip() != exam_client_ip:
        return jsonify({
            'success': False,
            'error': 'Nicht autorisiert',
            'message': 'Nur das Gerät, das die Prüfung gestartet hat, kann sie beenden.'
        }), 403

    exam_number = request.json.get('exam_number')
    duration = request.json.get('duration', 0)

    relay_controller.reset_all_relays()
    update_examination_duration(exam_number, duration)

    exam_active = False
    exam_client_ip = None
    exam_number_current = None
    exam_start_time = None
    print(f"🔓 Prüfung beendet")
    return jsonify({'success': True})


# ==================== MANUELLER MODUS ====================

@app.route('/manual_mode')
def manual_mode():
    """Manuelle Fehlerauswahl"""
    from managers.settings_manager import get_wallbox_enabled

    # Lade neue Relais-Konfiguration und dynamische Stromkreise
    relais_config = get_all_relais_config()
    groups = get_groups_overview()
    stromkreise = get_all_stromkreise()
    wallbox_enabled = get_wallbox_enabled()

    # Erstelle erweiterte Stromkreis-Info mit Relais
    stromkreise_with_names = {}

    for sk_id, sk_data in stromkreise.items():
        relay_options = []
        is_wallbox = sk_data['name'] == 'Wallbox'

        # Finde alle Relais mit diesem Stromkreis (Sicherheitsrelais ausschließen)
        relais_in_stromkreis = []
        for relay_num in range(64):
            if relay_num == SAFETY_RELAY_ID:
                continue
            relay_data = relais_config.get(relay_num, {})
            if relay_data.get('stromkreis') == sk_data['name']:
                relais_in_stromkreis.append(relay_num)

        # Gruppiere nach Gruppen-Nummer
        processed_groups = set()

        for relay_num in sorted(relais_in_stromkreis):
            relay_data = relais_config.get(relay_num, {})
            group_num = relay_data.get('group_number', 0)

            if group_num > 0:
                # Ist in Gruppe
                if group_num not in processed_groups:
                    processed_groups.add(group_num)
                    group_info = groups.get(group_num, {})
                    relay_options.append({
                        'number': relay_num,
                        'name': relay_data.get('name', f'Gruppe {group_num}'),
                        'is_group': True,
                        'group_relays': group_info.get('relays', [relay_num])
                    })
            else:
                # Einzelnes Relais
                relay_options.append({
                    'number': relay_num,
                    'name': relay_data.get('name', f'Relais {relay_num}'),
                    'is_group': False
                })

        stromkreise_with_names[sk_id] = {
            'name': sk_data['name'],
            'description': sk_data.get('description', ''),
            'relays': relay_options,
            'disabled': is_wallbox and not wallbox_enabled
        }

    return render_template('manual_mode_pi.html',
                         stromkreise=stromkreise_with_names,
                         wallbox_enabled=wallbox_enabled)


@app.route('/set_manual_errors', methods=['POST'])
def set_manual_errors():
    """Setzt manuell ausgewählte Fehler"""
    if exam_active:
        return exam_lock_response()
    try:
        data = request.json
        if not data:
            return jsonify({'success': False, 'error': 'Keine Daten erhalten'})
        
        errors = data.get('errors', {})
        if not errors:
            return jsonify({'success': False, 'error': 'Keine Fehler ausgewählt'})
        
        # Alle Relais zurücksetzen
        relay_controller.reset_all_relays()
        
        # Sammle nur eindeutige Repräsentanten (bei Gruppen)
        unique_relays = set()
        for stromkreis_key, relay_id in errors.items():
            try:
                relay_id = int(relay_id)
                if 0 <= relay_id <= 63:
                    # Normalisiere zu Gruppen-Repräsentant
                    representative = relay_controller.normalize_relay_to_group_representative(relay_id)
                    unique_relays.add(representative)
            except ValueError:
                pass
        
        activated_count = 0
        activated_relays = []
        failed_relays = []
        
        # Aktiviere eindeutige Relais/Gruppen
        for relay_id in unique_relays:
            if relay_controller.set_relay(relay_id, True):
                activated_count += 1
                activated_relays.append(relay_id)
                print(f"✅ Relay/Group {relay_id} activated")
            else:
                failed_relays.append(relay_id)
                print(f"❌ Failed to activate relay {relay_id}")
        
        return jsonify({
            'success': activated_count > 0,
            'activated_count': activated_count,
            'active_relays': activated_relays,
            'failed_relays': failed_relays,
            'message': f'{activated_count} Fehler zugeschaltet' + 
                      (f', {len(failed_relays)} fehlgeschlagen' if failed_relays else '')
        })
    
    except Exception as e:
        print(f"ERROR in set_manual_errors: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': f'Server error: {str(e)}'})


@app.route('/reset_relays', methods=['POST'])
def reset_relays():
    """Setzt alle Relais zurück"""
    if exam_active:
        return exam_lock_response()
    try:
        success = relay_controller.reset_all_relays()
        return jsonify({
            'success': success,
            'message': 'Alle Relais zurückgesetzt',
            'active_relays': relay_controller.active_relays
        })
    except Exception as e:
        print(f"ERROR in reset_relays: {e}")
        return jsonify({'success': False, 'error': str(e)})


@app.route('/relay_status')
def relay_status_page():
    """Relay Status Monitor Seite"""
    return render_template('relay_status.html')


@app.route('/api/relay_status', methods=['GET'])
def api_relay_status():
    """API Endpoint zum Auslesen aller Relay-Status"""
    try:
        relay_status = relay_controller.read_all_relay_status()

        if relay_status is None:
            return jsonify({
                'success': False,
                'error': 'Konnte Relais-Status nicht auslesen'
            })

        # Zähle aktive Relais
        active_count = sum(1 for state in relay_status.values() if state)
        active_relays = [num for num, state in relay_status.items() if state]

        return jsonify({
            'success': True,
            'relays': relay_status,
            'active_count': active_count,
            'total_count': len(relay_status),
            'active_relays': active_relays
        })
    except Exception as e:
        print(f"ERROR in api_relay_status: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        })


@app.route('/api/exam_status', methods=['GET'])
def api_exam_status():
    """Gibt den aktuellen Prüfungsstatus zurück"""
    import time as _time
    client_ip = get_client_ip()
    elapsed = int(_time.time() - exam_start_time) if exam_active and exam_start_time else 0
    return jsonify({
        'exam_active': exam_active,
        'is_exam_owner': exam_active and exam_client_ip == client_ip,
        'exam_client_ip': exam_client_ip if exam_active else None,
        'exam_number': exam_number_current if exam_active else None,
        'elapsed_seconds': elapsed
    })


# ==================== SICHERHEITSRELAIS ====================

@app.route('/api/safety_relay', methods=['GET'])
def api_safety_relay_status():
    """Gibt den aktuellen Sicherheitsrelais-Status zurück"""
    return jsonify({
        'enabled': relay_controller.safety_relay_state
    })


@app.route('/api/safety_relay/set', methods=['POST'])
def api_safety_relay_set():
    """Schaltet das Sicherheitsrelais. Einschalten erfordert Freigabe-Code."""
    from config import SAFETY_RELAY_ID
    from managers.settings_manager import verify_freigabe_code
    data = request.json or {}
    enable = data.get('enable', False)

    if enable:
        code = data.get('code', '')
        if not verify_freigabe_code(code):
            return jsonify({'success': False, 'error': 'Falscher Code'}), 403

    relay_controller.safety_relay_state = enable
    relay_controller.set_relay(SAFETY_RELAY_ID, enable)
    print(f"{'🟢' if enable else '🔴'} Sicherheitsrelais {'eingeschaltet' if enable else 'ausgeschaltet'}")
    return jsonify({'success': True, 'enabled': enable})


# ==================== TESTMODUS ====================

@app.route('/test_mode')
def test_mode():
    """Testmodus-Seite"""
    return render_template('test_mode.html')


@app.route('/run_test', methods=['POST'])
def run_test():
    """Führt einen vollständigen Relais-Test durch"""
    try:
        success = relay_controller.test_all_relays()
        return jsonify({'success': success})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/run_test_stream')
def run_test_stream():
    """Führt Relais-Test mit Live-Updates via Server-Sent Events durch"""
    def generate():
        try:
            yield f"data: {json.dumps({'type': 'start', 'total': 64})}\n\n"

            failed_relays = []

            for relay in range(64):
                # Relais einschalten
                yield f"data: {json.dumps({'type': 'testing', 'relay': relay, 'action': 'on'})}\n\n"

                if not relay_controller.set_relay(relay, True):
                    failed_relays.append({'relay': relay, 'error': 'Konnte nicht einschalten'})
                    yield f"data: {json.dumps({'type': 'error', 'relay': relay, 'error': 'Einschalten fehlgeschlagen'})}\n\n"
                    continue

                time.sleep(1.0)

                # Modbus-Readback prüfen ob Relais AN ist
                status = relay_controller.read_all_relay_status()
                modbus_ok = status and status.get(relay, False) == True

                yield f"data: {json.dumps({'type': 'status', 'relay': relay, 'modbus_ok': modbus_ok, 'state': 'on'})}\n\n"

                if not modbus_ok:
                    failed_relays.append({'relay': relay, 'error': 'Modbus: Relais nicht AN'})

                time.sleep(1.5)  # Grüne Anzeige länger sichtbar

                # Relais ausschalten
                yield f"data: {json.dumps({'type': 'testing', 'relay': relay, 'action': 'off'})}\n\n"

                if not relay_controller.set_relay(relay, False):
                    failed_relays.append({'relay': relay, 'error': 'Konnte nicht ausschalten'})

                time.sleep(0.5)

                # Modbus-Readback prüfen ob ALLE Relais AUS sind
                status = relay_controller.read_all_relay_status()
                any_on = False
                active_relays = []
                if status:
                    for r, state in status.items():
                        if state:
                            any_on = True
                            active_relays.append(r)

                all_off_ok = not any_on
                yield f"data: {json.dumps({'type': 'all_off', 'modbus_ok': all_off_ok, 'active_relays': active_relays})}\n\n"

                if not all_off_ok:
                    failed_relays.append({'relay': relay, 'error': f'Modbus: Relais noch aktiv: {active_relays}'})

                time.sleep(1.5)  # Blaue Anzeige länger sichtbar

                # Fortschritt
                yield f"data: {json.dumps({'type': 'progress', 'relay': relay, 'done': relay + 1, 'total': 64})}\n\n"

            # Test abgeschlossen
            yield f"data: {json.dumps({'type': 'complete', 'success': len(failed_relays) == 0, 'failed': failed_relays})}\n\n"

        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'error': str(e)})}\n\n"

    return Response(generate(), mimetype='text/event-stream')


@app.route('/test_single_relay/<int:relay_id>')
def test_single_relay(relay_id):
    """Testet ein einzelnes Relais"""
    if exam_active:
        return exam_lock_response()
    try:
        relay_controller.reset_all_relays()
        success = relay_controller.set_relay(relay_id, True)
        time.sleep(0.5)
        relay_controller.set_relay(relay_id, False)
        return jsonify({'success': success, 'relay_id': relay_id})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


# ==================== ÜBUNGSMODUS ====================

@app.route('/training_mode')
def training_mode():
    """Übungsmodus-Hauptseite"""
    return render_template('training_mode.html')


@app.route('/training/spannungsfrei')
def training_spannungsfrei():
    """Spannungsfreie Messungen (PE, RISO)"""
    return render_template('training_spannungsfrei.html')


@app.route('/training/unter_spannung')
def training_unter_spannung():
    """Messungen unter Spannung (Zs, Zi, RCD, Drehfeld)"""
    return render_template('training_unter_spannung.html')


# ==================== ADMIN-BEREICH ====================

@app.route('/admin')
def admin():
    """Admin-Login"""
    return render_template('admin_login.html')


@app.route('/admin_login', methods=['POST'])
def admin_login():
    """Admin-Login Validierung"""
    code = request.json.get('code')
    # Verwende den Code aus settings.json, falls vorhanden
    is_valid = verify_admin_code(code)
    return jsonify({'success': is_valid})


@app.route('/admin_panel')
def admin_panel():
    """Admin-Übersicht"""
    return render_template('admin_panel.html')





def reload_relay_config():
    """Lädt RELAY_GROUPS und RELAY_NAMES aus Dateien neu"""
    global RELAY_GROUPS
    try:
        # Update config mit neuen Daten
        import config
        config.RELAY_GROUPS = get_all_groups()
        config.RELAY_NAMES = get_all_relay_names()
        config.STROMKREISE = get_all_stromkreise()

        print("✓ Relay configuration reloaded")
        return True
    except Exception as e:
        print(f"Error reloading config: {e}")
        return False


# ==================== STROMKREIS-VERWALTUNG ====================


@app.route('/admin_config')
def admin_config():
    """Konfigurationsseite (Stromkreise + Kategorien in Tabs)"""
    stromkreise = get_all_stromkreise()
    kategorien = get_all_kategorien()

    return render_template('admin_config.html',
                         stromkreise=stromkreise,
                         kategorien=kategorien)


@app.route('/api/stromkreise', methods=['GET'])
def api_get_stromkreise():
    """API: Alle Stromkreise abrufen"""
    stats = get_stromkreis_statistics()
    return jsonify({
        'success': True,
        'stromkreise': stats['stromkreise'],
        'statistics': {
            'total_stromkreise': stats['total_stromkreise'],
            'unique_covered_relays': stats['unique_covered_relays'],
            'uncovered_relays': stats['uncovered_relays']
        }
    })


@app.route('/api/stromkreise/add', methods=['POST'])
def api_add_stromkreis():
    """API: Neuen Stromkreis erstellen"""
    data = request.json

    name = data.get('name', '').strip()
    description = data.get('description', '').strip()

    success, message, stromkreis_id = add_stromkreis(name, description)

    if success:
        reload_relay_config()

    return jsonify({
        'success': success,
        'message': message,
        'stromkreis_id': stromkreis_id
    })


@app.route('/api/stromkreise/update', methods=['POST'])
def api_update_stromkreis():
    """API: Stromkreis aktualisieren"""
    data = request.json

    stromkreis_id = data.get('stromkreis_id')
    name = data.get('name', '').strip()
    description = data.get('description', '').strip()

    success, message = update_stromkreis(stromkreis_id, name, description)

    if success:
        reload_relay_config()

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/stromkreise/delete', methods=['POST'])
def api_delete_stromkreis():
    """API: Stromkreis löschen"""
    data = request.json
    stromkreis_id = data.get('stromkreis_id')

    success, message = delete_stromkreis(stromkreis_id)

    if success:
        reload_relay_config()

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/kategorien', methods=['GET'])
def api_get_kategorien():
    """API: Alle Kategorien abrufen"""
    kategorien = get_all_kategorien()
    return jsonify({
        'success': True,
        'kategorien': kategorien
    })


@app.route('/api/kategorien/add', methods=['POST'])
def api_add_kategorie():
    """API: Neue Kategorie hinzufügen"""
    data = request.json
    name = data.get('name', '').strip()

    success, message = add_kategorie(name)

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/kategorien/delete', methods=['POST'])
def api_delete_kategorie():
    """API: Kategorie löschen"""
    data = request.json
    name = data.get('name', '')

    success, message = delete_kategorie(name)

    return jsonify({
        'success': success,
        'message': message
    })


# ==================== RELAIS-VERWALTUNG (NEUE MODULARE VERSION) ====================

@app.route('/admin_relais')
def admin_relais():
    """Neue modulare Relais-Verwaltungsseite"""
    from managers.settings_manager import get_wallbox_installed

    relais_config = get_all_relais_config()
    stats = get_relais_statistics()
    kategorien = get_all_kategorien()
    stromkreise = get_all_stromkreise()
    wallbox_installed = get_wallbox_installed()

    return render_template('admin_relais.html',
                         relais_config=relais_config,
                         stats=stats,
                         kategorien=kategorien,
                         stromkreise=stromkreise,
                         wallbox_installed=wallbox_installed)


@app.route('/api/relais/config', methods=['GET'])
def api_get_relais_config():
    """API: Alle Relais-Konfigurationen abrufen"""
    config = get_all_relais_config()
    return jsonify({
        'success': True,
        'relais_config': config
    })


@app.route('/api/relais/statistics', methods=['GET'])
def api_get_relais_statistics():
    """API: Statistiken über Relais-Konfiguration"""
    stats = get_relais_statistics()
    return jsonify({
        'success': True,
        'statistics': stats
    })


@app.route('/api/relais/update', methods=['POST'])
def api_update_relais():
    """API: Einzelnes Relais aktualisieren"""
    if exam_active:
        return exam_lock_response()
    data = request.json

    relay_num = data.get('relay_num')
    group_number = data.get('group_number', 0)
    name = data.get('name', '')
    category = data.get('category', '')
    stromkreis = data.get('stromkreis', '')

    try:
        relay_num = int(relay_num)
        group_number = int(group_number)
    except:
        return jsonify({'success': False, 'message': 'Ungültige Eingabe'})

    success, message = update_relay_config(relay_num, group_number, name, category, stromkreis)

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/relais/bulk_update', methods=['POST'])
def api_bulk_update_relais():
    """API: Mehrere Relais auf einmal aktualisieren"""
    if exam_active:
        return exam_lock_response()
    data = request.json
    updates = data.get('updates', {})

    result = bulk_update_relais(updates)

    return jsonify({
        'success': result['success'],
        'message': result['message'],
        'failed_count': result['failed_count']
    })


@app.route('/api/relais/groups', methods=['GET'])
def api_get_relais_groups():
    """API: Übersicht über alle Gruppen"""
    groups = get_groups_overview()

    return jsonify({
        'success': True,
        'groups': groups
    })


@app.route('/api/relais/by_category/<category>', methods=['GET'])
def api_get_relais_by_category_new(category):
    """API: Relais einer Kategorie abrufen"""
    relais_list = get_relais_by_category(category)

    return jsonify({
        'success': True,
        'category': category,
        'relais': relais_list
    })


# ==================== RELAIS-VORLAGEN ====================

@app.route('/api/relais/templates', methods=['GET'])
def api_get_templates():
    """API: Verfügbare Vorlagen abrufen"""
    templates = get_available_templates()

    return jsonify({
        'success': True,
        'templates': templates
    })


@app.route('/api/relais/template/<template_id>', methods=['GET'])
def api_get_template_info(template_id):
    """API: Detaillierte Informationen zu einer Vorlage"""
    template_info = get_template_info(template_id)

    if template_info:
        return jsonify({
            'success': True,
            'template': template_info
        })
    else:
        return jsonify({
            'success': False,
            'message': 'Vorlage nicht gefunden'
        }), 404


@app.route('/api/relais/apply_template', methods=['POST'])
def api_apply_template():
    """API: Vorlage anwenden"""
    if exam_active:
        return exam_lock_response()
    try:
        data = request.json
        template_id = data.get('template_id')
        start_relay = data.get('start_relay')
        group_offset = data.get('group_offset', 0)

        if not template_id:
            return jsonify({
                'success': False,
                'message': 'Keine Vorlage ausgewählt'
            }), 400

        # Vorlage anwenden
        template_config = apply_template(template_id, start_relay, group_offset)

        if not template_config:
            return jsonify({
                'success': False,
                'message': 'Vorlage konnte nicht geladen werden'
            }), 404

        # Bulk-Update durchführen
        result = bulk_update_relais(template_config)

        if result['success']:
            return jsonify({
                'success': True,
                'message': f"Vorlage '{template_id}' erfolgreich angewendet",
                'applied_relais': list(template_config.keys()),
                'count': len(template_config)
            })
        else:
            return jsonify({
                'success': False,
                'message': result.get('message', 'Fehler beim Anwenden der Vorlage')
            }), 500

    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        }), 500


# ==================== EXCEL IMPORT/EXPORT ====================

@app.route('/api/relais/excel/template/empty', methods=['GET'])
def api_download_empty_template():
    """API: Leere Excel-Vorlage herunterladen"""
    try:
        excel_file = create_excel_template(include_current_config=False)

        return send_file(
            excel_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='VDE_Messwand_Relais_Vorlage_Leer.xlsx'
        )
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler beim Erstellen der Vorlage: {str(e)}'
        }), 500


@app.route('/api/relais/excel/template/current', methods=['GET'])
def api_download_current_config():
    """API: Excel-Vorlage mit aktueller Konfiguration herunterladen"""
    try:
        excel_file = create_excel_template(include_current_config=True)

        return send_file(
            excel_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='VDE_Messwand_Relais_Aktuell.xlsx'
        )
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler beim Erstellen der Datei: {str(e)}'
        }), 500


@app.route('/api/relais/excel/template/predefined', methods=['GET'])
def api_get_predefined_templates():
    """API: Liste vordefinierter Excel-Vorlagen"""
    templates = get_predefined_templates()

    # Nur Metadaten zurückgeben, keine Config
    template_list = [{
        'id': t['id'],
        'name': t['name'],
        'description': t['description']
    } for t in templates]

    return jsonify({
        'success': True,
        'templates': template_list
    })


@app.route('/api/relais/excel/template/predefined/<template_id>', methods=['GET'])
def api_download_predefined_template(template_id):
    """API: Vordefinierte Excel-Vorlage herunterladen"""
    try:
        excel_file = create_predefined_template_excel(template_id)

        if not excel_file:
            return jsonify({
                'success': False,
                'message': 'Vorlage nicht gefunden'
            }), 404

        return send_file(
            excel_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'VDE_Messwand_{template_id}.xlsx'
        )
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler beim Erstellen der Vorlage: {str(e)}'
        }), 500


@app.route('/api/relais/excel/import', methods=['POST'])
def api_import_excel():
    """API: Excel-Datei importieren"""
    if exam_active:
        return exam_lock_response()
    try:
        if 'file' not in request.files:
            return jsonify({
                'success': False,
                'message': 'Keine Datei hochgeladen'
            }), 400

        file = request.files['file']

        if file.filename == '':
            return jsonify({
                'success': False,
                'message': 'Keine Datei ausgewählt'
            }), 400

        if not file.filename.endswith(('.xlsx', '.xls')):
            return jsonify({
                'success': False,
                'message': 'Nur Excel-Dateien (.xlsx, .xls) sind erlaubt'
            }), 400

        # Importiere Daten
        result = import_from_excel(file.stream)

        if result['success']:
            return jsonify(result)
        else:
            return jsonify(result), 400

    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler beim Importieren: {str(e)}'
        }), 500


# ==================== ÜBUNGSMODUS-VERWALTUNG ====================

@app.route('/admin_training')
def admin_training():
    """Übungsmodus-Admin-Seite - NEUE Kategorie-basierte Ansicht"""
    training_config = get_complete_training_config()
    training_pages = get_training_pages()
    kategorien = get_all_kategorien()
    stats = get_statistics()
    relais_config = get_all_relais_config()

    return render_template('admin_training_new.html',
                         training_config=training_config,
                         training_pages=training_pages,
                         kategorien=kategorien,
                         stats=stats,
                         relais_config=relais_config)


@app.route('/api/training/config', methods=['GET'])
def api_get_training_config():
    """API: Komplette Training-Konfiguration abrufen"""
    config = get_complete_training_config()

    return jsonify({
        'success': True,
        'training_config': config
    })


@app.route('/api/training/pages', methods=['GET'])
def api_get_training_pages():
    """API: Verfügbare Übungsseiten"""
    pages = get_training_pages()

    return jsonify({
        'success': True,
        'pages': pages
    })


@app.route('/api/training/<page_id>/<category>', methods=['GET'])
def api_get_training_relais(page_id, category):
    """API: Relais für Übungsseite/Kategorie abrufen"""
    relais_list = get_relais_for_training(page_id, category)

    return jsonify({
        'success': True,
        'page_id': page_id,
        'category': category,
        'relais': relais_list
    })


@app.route('/api/training/update', methods=['POST'])
def api_update_training_mapping():
    """API: Mapping für Kategorie/Übungsseite aktualisieren"""
    data = request.json

    category = data.get('category', '')
    page_id = data.get('page_id', '')
    relais_list = data.get('relais_list', [])

    success, message = update_training_mapping(category, page_id, relais_list)

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/training/delete', methods=['POST'])
def api_delete_training_mapping():
    """API: Mapping löschen"""
    data = request.json

    category = data.get('category', '')
    page_id = data.get('page_id', '')

    success, message = delete_training_mapping(category, page_id)

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/training/import_from_category', methods=['POST'])
def api_import_training_from_category():
    """API: Auto-Import aus Relais-Manager"""
    data = request.json

    category = data.get('category', '')
    page_id = data.get('page_id', '')

    success, message, count = import_from_relais_manager(category, page_id)

    return jsonify({
        'success': success,
        'message': message,
        'imported_count': count
    })


# ==================== WARTUNGSTEST ====================

@app.route('/wartungstest')
def wartungstest():
    """Wartungstest-Seite: alle Relais nacheinander testen"""
    import json as _json
    from managers.relais_manager import load_relais_config
    config = load_relais_config()
    # Sortiert nach Relais-Nummer
    relais = sorted(
        [{'id': int(k), 'name': v.get('name', ''), 'category': v.get('category', ''), 'stromkreis': v.get('stromkreis', '')}
         for k, v in config.items()],
        key=lambda x: x['id']
    )
    # Stromkreise aus den tatsächlichen Relais-Daten ableiten (modular)
    sk_counts = {}
    for r in relais:
        sk = r['stromkreis']
        if sk:
            sk_counts[sk] = sk_counts.get(sk, 0) + 1
    stromkreise = [{'name': k, 'count': v} for k, v in sorted(sk_counts.items())]

    # Kategorien aus den tatsächlichen Relais-Daten ableiten
    kat_counts = {}
    for r in relais:
        kat = r['category']
        if kat:
            kat_counts[kat] = kat_counts.get(kat, 0) + 1
    kategorien = [{'name': k, 'count': v} for k, v in sorted(kat_counts.items())]

    return render_template('wartungstest.html',
                           relais=relais,
                           relais_json=_json.dumps(relais),
                           stromkreise=stromkreise,
                           kategorien=kategorien)


@app.route('/api/wartungstest/relay/<int:relay_id>', methods=['POST'])
def api_wartungstest_relay(relay_id):
    """Schaltet exklusiv ein Relais für den Wartungstest"""
    if exam_active:
        return exam_lock_response()
    try:
        relay_controller.reset_all_relays()
        success = relay_controller.set_relay(relay_id, True)
        return jsonify({'success': success, 'relay_id': relay_id})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/wartungstest/export/pdf', methods=['POST'])
def api_wartungstest_export_pdf():
    """Exportiert Testergebnisse als PDF und speichert sie im pdfs/-Ordner"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT

        data = request.json or {}
        results = data.get('results', [])
        tester = data.get('tester', '')
        datum = data.get('datum', datetime.now().strftime('%d.%m.%Y'))

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
                                leftMargin=1.5*cm, rightMargin=1.5*cm,
                                topMargin=1.5*cm, bottomMargin=1.5*cm)

        styles = getSampleStyleSheet()
        red = colors.HexColor('#E30613')
        green = colors.HexColor('#27AE60')
        dark_red = colors.HexColor('#C0392B')
        grey = colors.HexColor('#7F8C8D')
        dark_bg = colors.HexColor('#1A1A2E')

        title_style = ParagraphStyle('title', fontSize=18, textColor=colors.white,
                                     backColor=red, alignment=TA_CENTER,
                                     spaceAfter=4, spaceBefore=4,
                                     leftIndent=-10, rightIndent=-10)
        meta_style = ParagraphStyle('meta', fontSize=10, textColor=colors.HexColor('#555555'))

        story = []

        # Titel
        story.append(Paragraph('VDE Messwand – Wartungstest', title_style))
        story.append(Spacer(1, 0.3*cm))

        meta_text = f'Datum: {datum}'
        if tester:
            meta_text += f'   |   Prüfer: {tester}'
        story.append(Paragraph(meta_text, meta_style))
        story.append(Spacer(1, 0.4*cm))

        # Tabelle
        header = ['Nr.', 'Name', 'Kategorie', 'Stromkreis', 'Ergebnis', 'Kommentar']
        table_data = [header]
        row_colors = []

        for item in results:
            status = item.get('status')
            ergebnis = 'OK ✔' if status == 'ok' else ('FEHLER ✘' if status == 'fail' else '–')
            table_data.append([
                str(item.get('id', 0) + 1),
                item.get('name', ''),
                item.get('category', ''),
                item.get('stromkreis', ''),
                ergebnis,
                item.get('kommentar', '')
            ])
            row_colors.append(status)

        col_widths = [1.2*cm, 5.5*cm, 3*cm, 3.5*cm, 2.5*cm, 3.3*cm]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)

        ts = TableStyle([
            # Header
            ('BACKGROUND', (0,0), (-1,0), dark_bg),
            ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
            ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,0), 9),
            ('ALIGN',      (0,0), (-1,0), 'CENTER'),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
            ('TOPPADDING',    (0,0), (-1,0), 6),
            # Body
            ('FONTNAME',   (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE',   (0,1), (-1,-1), 8.5),
            ('VALIGN',     (0,0), (-1,-1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#F8F8F8'), colors.white]),
            ('GRID',       (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
            ('TOPPADDING',    (0,1), (-1,-1), 4),
            ('BOTTOMPADDING', (0,1), (-1,-1), 4),
        ])

        # Ergebnis-Spalte einfärben
        for i, status in enumerate(row_colors, 1):
            if status == 'ok':
                ts.add('BACKGROUND', (4,i), (4,i), green)
                ts.add('TEXTCOLOR',  (4,i), (4,i), colors.white)
                ts.add('FONTNAME',   (4,i), (4,i), 'Helvetica-Bold')
            elif status == 'fail':
                ts.add('BACKGROUND', (4,i), (4,i), dark_red)
                ts.add('TEXTCOLOR',  (4,i), (4,i), colors.white)
                ts.add('FONTNAME',   (4,i), (4,i), 'Helvetica-Bold')
            else:
                ts.add('BACKGROUND', (4,i), (4,i), grey)
                ts.add('TEXTCOLOR',  (4,i), (4,i), colors.white)

        t.setStyle(ts)
        story.append(t)
        story.append(Spacer(1, 0.5*cm))

        # Statistik
        total = len(results)
        ok_count   = sum(1 for r in results if r.get('status') == 'ok')
        fail_count = sum(1 for r in results if r.get('status') == 'fail')
        skip_count = total - ok_count - fail_count

        stat_data = [['Gesamt', 'In Ordnung', 'Fehler', 'Übersprungen'],
                     [str(total), str(ok_count), str(fail_count), str(skip_count)]]
        st = Table(stat_data, colWidths=[3*cm]*4)
        st.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), dark_bg),
            ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
            ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE',   (0,0), (-1,0), 9),
            ('ALIGN',      (0,0), (-1,-1), 'CENTER'),
            ('FONTSIZE',   (0,1), (-1,-1), 14),
            ('FONTNAME',   (0,1), (-1,-1), 'Helvetica-Bold'),
            ('TEXTCOLOR',  (0,1), (0,1), colors.black),
            ('TEXTCOLOR',  (1,1), (1,1), green),
            ('TEXTCOLOR',  (2,1), (2,1), dark_red),
            ('TEXTCOLOR',  (3,1), (3,1), grey),
            ('TOPPADDING',    (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CCCCCC')),
        ]))
        story.append(st)

        doc.build(story)
        buf.seek(0)

        filename = f"Wartungstest_{datum.replace('.', '-')}.pdf"

        # Im pdfs/-Ordner speichern
        import os as _os
        if not _os.path.exists(PDF_DIR):
            _os.makedirs(PDF_DIR)
        save_path = _os.path.join(PDF_DIR, filename)
        with open(save_path, 'wb') as f:
            f.write(buf.getvalue())
        buf.seek(0)

        return send_file(buf, as_attachment=True, download_name=filename,
                         mimetype='application/pdf')
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/wartungstest/export', methods=['POST'])
def api_wartungstest_export():
    """Exportiert Testergebnisse als Excel-Datei"""
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        data = request.json or {}
        results = data.get('results', [])
        tester = data.get('tester', '')
        datum = data.get('datum', datetime.now().strftime('%d.%m.%Y'))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Wartungstest'

        # Farben
        red_fill   = PatternFill('solid', fgColor='C0392B')
        green_fill = PatternFill('solid', fgColor='27AE60')
        grey_fill  = PatternFill('solid', fgColor='7F8C8D')
        head_fill  = PatternFill('solid', fgColor='1A1A2E')
        thin_border = Border(
            left=Side(style='thin', color='AAAAAA'),
            right=Side(style='thin', color='AAAAAA'),
            top=Side(style='thin', color='AAAAAA'),
            bottom=Side(style='thin', color='AAAAAA')
        )

        # Titel-Zeile
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = 'VDE Messwand – Wartungstest'
        title_cell.font = Font(name='Calibri', bold=True, size=16, color='FFFFFF')
        title_cell.fill = PatternFill('solid', fgColor='E30613')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 32

        # Meta-Zeile
        ws.merge_cells('A2:C2')
        ws['A2'].value = f'Datum: {datum}'
        ws['A2'].font = Font(name='Calibri', size=11)
        ws.merge_cells('D2:F2')
        ws['D2'].value = f'Prüfer: {tester}' if tester else ''
        ws['D2'].font = Font(name='Calibri', size=11)
        ws.row_dimensions[2].height = 20

        # Header-Zeile
        headers = ['Relais-Nr.', 'Name', 'Kategorie', 'Stromkreis', 'Ergebnis', 'Kommentar']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.fill = head_fill
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        ws.row_dimensions[3].height = 22

        # Daten
        for row_idx, item in enumerate(results, 4):
            status = item.get('status')
            ergebnis = 'OK ✔' if status == 'ok' else ('FEHLER ✘' if status == 'fail' else 'Übersprungen')
            row_data = [
                (item.get('id', 0) + 1),  # 1-basierte Anzeige
                item.get('name', ''),
                item.get('category', ''),
                item.get('stromkreis', ''),
                ergebnis,
                item.get('kommentar', '')
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.font = Font(name='Calibri', size=10)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = thin_border
                if col == 5:
                    if status == 'ok':
                        cell.fill = green_fill
                        cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
                    elif status == 'fail':
                        cell.fill = red_fill
                        cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
                    else:
                        cell.fill = grey_fill
                        cell.font = Font(name='Calibri', size=10, color='FFFFFF')
            ws.row_dimensions[row_idx].height = 18

        # Spaltenbreiten
        col_widths = [12, 40, 20, 25, 16, 40]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # Statistik-Block
        total = len(results)
        ok_count   = sum(1 for r in results if r.get('status') == 'ok')
        fail_count = sum(1 for r in results if r.get('status') == 'fail')
        skip_count = total - ok_count - fail_count
        stat_row = total + 5
        ws.cell(row=stat_row, column=1, value='Gesamt:').font = Font(bold=True)
        ws.cell(row=stat_row, column=2, value=total)
        ws.cell(row=stat_row+1, column=1, value='In Ordnung:').font = Font(bold=True, color='27AE60')
        ws.cell(row=stat_row+1, column=2, value=ok_count)
        ws.cell(row=stat_row+2, column=1, value='Fehler:').font = Font(bold=True, color='C0392B')
        ws.cell(row=stat_row+2, column=2, value=fail_count)
        ws.cell(row=stat_row+3, column=1, value='Übersprungen:').font = Font(bold=True, color='7F8C8D')
        ws.cell(row=stat_row+3, column=2, value=skip_count)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = f"Wartungstest_{datum.replace('.', '-')}.xlsx"

        # Zusätzlich im pdfs/-Ordner speichern (für Dokument-Anzeige im System)
        try:
            import os as _os
            if not _os.path.exists(PDF_DIR):
                _os.makedirs(PDF_DIR)
            save_path = _os.path.join(PDF_DIR, filename)
            with open(save_path, 'wb') as f:
                f.write(buf.getvalue())
            buf.seek(0)
        except Exception:
            buf.seek(0)

        return send_file(buf, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/training/statistics', methods=['GET'])
def api_get_training_statistics():
    """API: Statistiken über Training-Config"""
    stats = get_statistics()

    return jsonify({
        'success': True,
        'statistics': stats
    })


@app.route('/api/training/activate', methods=['POST'])
def api_activate_training():
    """API: Aktiviert Relais für Übungsseite/Kategorie"""
    if exam_active:
        return exam_lock_response()
    data = request.json

    page_id = data.get('page_id', '')
    category = data.get('category', '')

    # Hole konfigurierte Relais
    relais_list = get_relais_for_training(page_id, category)

    if not relais_list:
        return jsonify({
            'success': False,
            'message': f'Keine Relais für {page_id}/{category} konfiguriert'
        })

    try:
        # Alle Relais zurücksetzen
        relay_controller.reset_all_relays()

        # Aktiviere konfigurierte Relais (mit Gruppen-Normalisierung)
        activated = []
        for relay_num in relais_list:
            # Normalisiere zu Gruppen-Repräsentant
            representative = normalize_relay_to_representative(relay_num)

            if relay_controller.set_relay(representative, True):
                activated.append(relay_num)

        return jsonify({
            'success': True,
            'message': f'{len(activated)} Relais aktiviert',
            'activated_relais': activated
        })

    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        })


@app.route('/api/training/relays_by_category/<category>', methods=['GET'])
def api_training_relays_by_category(category):
    """API: Alle Relais einer Kategorie mit Details (ohne page_id-Filter)"""
    from managers.relais_manager import load_relais_config, get_relais_by_category
    relay_nums = get_relais_by_category(category)
    relais_config = load_relais_config()
    relays = []
    for relay_num in relay_nums:
        details = relais_config.get(str(relay_num), {})
        relays.append({
            'relay_num': relay_num,
            'name': details.get('name', f'Relais {relay_num}'),
            'stromkreis': details.get('stromkreis', ''),
        })
    return jsonify({'success': True, 'relays': relays})


@app.route('/api/training/relay_details/<page_id>/<category>', methods=['GET'])
def api_training_relay_details(page_id, category):
    """API: Relais mit Details für Übungsseite/Kategorie (für Einzel-Buttons)"""
    relays = get_training_relays_with_details(page_id, category)
    return jsonify({'success': True, 'relays': relays})


@app.route('/api/training/activate_single', methods=['POST'])
def api_training_activate_single():
    """API: Einzelnes Relais im Übungsmodus aktivieren"""
    if exam_active:
        return exam_lock_response()
    data = request.json or {}
    relay_num = data.get('relay_num')
    if relay_num is None:
        return jsonify({'success': False, 'message': 'relay_num fehlt'})
    try:
        representative = normalize_relay_to_representative(int(relay_num))
        success = relay_controller.set_relay(representative, True)
        return jsonify({
            'success': success,
            'message': f'Relais {relay_num} aktiviert' if success else f'Relais {relay_num} konnte nicht aktiviert werden'
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})


@app.route('/api/training/deactivate_single', methods=['POST'])
def api_training_deactivate_single():
    """API: Einzelnes Relais im Übungsmodus deaktivieren"""
    if exam_active:
        return exam_lock_response()
    data = request.json or {}
    relay_num = data.get('relay_num')
    if relay_num is None:
        return jsonify({'success': False, 'message': 'relay_num fehlt'})
    try:
        representative = normalize_relay_to_representative(int(relay_num))
        success = relay_controller.set_relay(representative, False)
        return jsonify({
            'success': success,
            'message': f'Relais {relay_num} deaktiviert' if success else f'Relais {relay_num} konnte nicht deaktiviert werden'
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)})


@app.route('/admin_database')
def admin_database():
    """Datenbank-Verwaltung"""
    import re as _re
    from managers.group_manager import get_all_groups
    from managers.relais_manager import get_all_relais_config

    examinations = get_all_examinations()
    stats = get_examination_stats()
    groups = get_all_groups()
    relais_config = get_all_relais_config()  # {0: {name, category, ...}, 1: ..., }

    def resolve_name(relay_num):
        """Gibt den konfigurierten Fehlernamen zurück, Fallback 'Relais X'."""
        cfg = relais_config.get(relay_num, {})
        name = cfg.get('name', '').strip()
        return name if name else f'Relais {relay_num}'

    # Formatiere Daten für Template mit Namen
    formatted_exams = []
    for exam in examinations:
        formatted_date, formatted_time = format_timestamp(exam['timestamp'])

        relay_descriptions = []
        for relay_entry in exam['active_relays']:
            if isinstance(relay_entry, int):
                relay_num = relay_entry
            else:
                # String-Eintrag: entweder echter Name oder Fallback "Relais X"
                stored_name = str(relay_entry)
                m = _re.fullmatch(r'Relais\s+(\d+)', stored_name)
                if m:
                    relay_num = int(m.group(1))
                else:
                    # Prüfe ob es ein gespeicherter Gruppenname ist (Altdaten)
                    matched_group = next(
                        ((gid, gd) for gid, gd in groups.items() if gd.get('name') == stored_name),
                        None
                    )
                    if matched_group:
                        _, gd = matched_group
                        # Repräsentanten-Relais der Gruppe holen und dessen Namen anzeigen
                        group_relays = gd.get('relays', [])
                        rep_relay = min(group_relays) if group_relays else None
                        if rep_relay is not None:
                            rep_name = resolve_name(rep_relay)
                            group_category = relais_config.get(rep_relay, {}).get('category', '') or gd.get('category', '')
                        else:
                            rep_name = stored_name
                            group_category = gd.get('category', '')
                        relay_descriptions.append({
                            'number': rep_relay,
                            'name': rep_name,
                            'is_group': True,
                            'category': group_category
                        })
                    else:
                        # Echter gespeicherter Relais-Name – direkt verwenden
                        relay_descriptions.append({'number': None, 'name': stored_name, 'is_group': False, 'category': ''})
                    continue

            name = resolve_name(relay_num)
            is_group = False
            group_category = ''
            for gd in groups.values():
                if relay_num in gd['relays']:
                    is_group = True
                    group_category = gd.get('category', '')
                    break
            relay_descriptions.append({
                'number': relay_num,
                'name': name,
                'is_group': is_group,
                'category': group_category
            })

        formatted_exams.append({
            'id': exam['id'],
            'exam_number': exam['exam_number'],
            'relay_list': exam['active_relays'],
            'relay_descriptions': relay_descriptions,
            'formatted_date': formatted_date,
            'formatted_time': formatted_time,
            'formatted_duration': format_duration(exam['duration']),
            'is_completed': exam['is_completed']
        })

    return render_template('admin_database.html',
                         examinations=formatted_exams,
                         completed_count=stats['completed'],
                         incomplete_count=stats['incomplete'])


@app.route('/db')
def database_display():
    """Desktop-Anzeige der Prüfungsdatenbank (Direktlink)"""
    import re as _re
    from managers.group_manager import get_all_groups
    from managers.relais_manager import get_all_relais_config

    examinations = get_all_examinations()
    stats = get_examination_stats()
    groups = get_all_groups()
    relais_config = get_all_relais_config()

    def resolve_name(relay_num):
        cfg = relais_config.get(relay_num, {})
        name = cfg.get('name', '').strip()
        return name if name else f'Relais {relay_num}'

    formatted_exams = []
    for exam in examinations:
        formatted_date, formatted_time = format_timestamp(exam['timestamp'])

        relay_descriptions = []
        for relay_entry in exam['active_relays']:
            if isinstance(relay_entry, int):
                relay_num = relay_entry
            else:
                stored_name = str(relay_entry)
                m = _re.fullmatch(r'Relais\s+(\d+)', stored_name)
                if m:
                    relay_num = int(m.group(1))
                else:
                    matched_group = next(
                        ((gid, gd) for gid, gd in groups.items() if gd.get('name') == stored_name),
                        None
                    )
                    if matched_group:
                        _, gd = matched_group
                        group_relays = gd.get('relays', [])
                        rep_relay = min(group_relays) if group_relays else None
                        if rep_relay is not None:
                            rep_name = resolve_name(rep_relay)
                            group_category = relais_config.get(rep_relay, {}).get('category', '') or gd.get('category', '')
                        else:
                            rep_name = stored_name
                            group_category = gd.get('category', '')
                        relay_descriptions.append({'number': rep_relay, 'name': rep_name, 'is_group': True, 'category': group_category})
                    else:
                        relay_descriptions.append({'number': None, 'name': stored_name, 'is_group': False, 'category': ''})
                    continue

            name = resolve_name(relay_num)
            is_group = False
            group_category = ''
            for gd in groups.values():
                if relay_num in gd['relays']:
                    is_group = True
                    group_category = gd.get('category', '')
                    break
            relay_descriptions.append({'number': relay_num, 'name': name, 'is_group': is_group, 'category': group_category})

        formatted_exams.append({
            'id': exam['id'],
            'exam_number': exam['exam_number'],
            'relay_descriptions': relay_descriptions,
            'formatted_date': formatted_date,
            'formatted_time': formatted_time,
            'formatted_duration': format_duration(exam['duration']),
            'is_completed': exam['is_completed']
        })

    avg_min = int(stats['avg_duration'] // 60) if stats['avg_duration'] else 0
    avg_sec = int(stats['avg_duration'] % 60) if stats['avg_duration'] else 0
    avg_duration_str = f"{avg_min}:{avg_sec:02d}" if stats['avg_duration'] else "—"

    return render_template('database_display.html',
                           examinations=formatted_exams,
                           total_count=stats['total'],
                           completed_count=stats['completed'],
                           incomplete_count=stats['incomplete'],
                           avg_duration_str=avg_duration_str)


@app.route('/clear_database', methods=['POST'])
def clear_database_route():
    """Löscht alle Prüfungsdaten"""
    success = clear_database()
    return jsonify({'success': success})


@app.route('/export_database')
def export_database():
    """Exportiert Datenbank als CSV"""
    try:
        csv_data = export_to_csv()
        
        output = io.StringIO()
        writer = csv.writer(output, delimiter=';')
        writer.writerows(csv_data)
        
        csv_content = output.getvalue()
        output.close()
        
        filename = f"vde_pruefungen_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        
        return Response(csv_content, mimetype='text/csv',
                       headers={'Content-Disposition': f'attachment; filename={filename}'})
    
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/admin_network')
def admin_network():
    """Netzwerk-Informationen"""
    try:
        ifconfig_result = subprocess.run(['ifconfig'], capture_output=True, text=True)
        network_info = ifconfig_result.stdout
        hostname = subprocess.run(['hostname'], capture_output=True, text=True).stdout.strip()

        # Hotspot-Status
        hotspot_active = is_hotspot_active()
        current_connection = get_current_connection()
        net_info = get_network_info()
        eth_info = get_ethernet_info()

        return render_template('admin_network.html',
                             network_info=network_info,
                             hostname=hostname,
                             hotspot_active=hotspot_active,
                             current_connection=current_connection,
                             net_info=net_info,
                             eth_info=eth_info)
    except Exception as e:
        return render_template('admin_network.html',
                             network_info=f"Error: {e}",
                             hostname="Unknown",
                             hotspot_active=False,
                             current_connection=[],
                             net_info={},
                             eth_info=None)


@app.route('/api/network/hotspot/toggle', methods=['POST'])
def api_toggle_hotspot():
    """API: WiFi-Hotspot ein/ausschalten"""
    try:
        success, message = toggle_hotspot()
        return jsonify({
            'success': success,
            'message': message,
            'hotspot_active': is_hotspot_active()
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        })


@app.route('/api/network/wifi/scan', methods=['GET'])
def api_scan_wifi():
    """API: WiFi-Netzwerke scannen"""
    try:
        networks = get_wifi_networks()
        return jsonify({
            'success': True,
            'networks': networks
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        })


@app.route('/api/network/wifi/connect', methods=['POST'])
def api_connect_wifi():
    """API: Mit WiFi-Netzwerk verbinden"""
    try:
        data = request.json
        ssid = data.get('ssid', '')
        password = data.get('password', '')

        if not ssid:
            return jsonify({
                'success': False,
                'message': 'SSID fehlt'
            })

        success, message = connect_to_wifi(ssid, password)
        return jsonify({
            'success': success,
            'message': message
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        })


@app.route('/api/network/status', methods=['GET'])
def api_network_status():
    """API: Netzwerk-Status abrufen"""
    try:
        hotspot_active = is_hotspot_active()
        current_connection = get_current_connection()
        net_info = get_network_info()

        return jsonify({
            'success': True,
            'hotspot_active': hotspot_active,
            'current_connection': current_connection,
            'network_info': net_info
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        })


# ==================== GPIO-MONITOR (SCHLIESSEN) ====================

@app.route('/api/gpio/status', methods=['GET'])
def api_gpio_status():
    """API: GPIO-Schließer-Status abrufen (ohne Logging)"""
    try:
        status = get_gpio_status()
        return jsonify({
            'success': True,
            'gpio_status': status
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'message': f'Fehler: {str(e)}'
        })


@app.route('/admin_settings')
def admin_settings():
    """Admin-Einstellungen"""
    from managers.settings_manager import get_gpio_shutdown_timeout, get_freigabe_code
    current_code = get_admin_code()
    masked_code = '*' * len(current_code)
    freigabe_code = get_freigabe_code()
    masked_freigabe = '*' * len(freigabe_code)
    gpio_timeout = get_gpio_shutdown_timeout()
    return render_template('admin_settings.html',
                           current_code=masked_code,
                           freigabe_code=masked_freigabe,
                           gpio_timeout=gpio_timeout)


@app.route('/api/settings/gpio_timeout', methods=['POST'])
def api_save_gpio_timeout():
    """API: Notaus-Shutdown-Timeout speichern"""
    from managers.settings_manager import get_gpio_shutdown_timeout, set_gpio_shutdown_timeout
    data = request.json or {}
    try:
        seconds = int(data.get('seconds', 120))
    except (ValueError, TypeError):
        return jsonify({'success': False, 'message': 'Ungültiger Wert'}), 400
    success, message = set_gpio_shutdown_timeout(seconds)
    if success:
        update_gpio_shutdown_timeout(seconds)
    return jsonify({'success': success, 'message': message})


@app.route('/api/settings/change_code', methods=['POST'])
def api_change_admin_code():
    """API: Admin-Code ändern"""
    data = request.json
    current_code = data.get('current_code', '')
    new_code = data.get('new_code', '')

    # Überprüfe aktuellen Code
    if not verify_admin_code(current_code):
        return jsonify({
            'success': False,
            'message': 'Aktueller Code ist falsch'
        })

    # Setze neuen Code
    success, message = set_admin_code(new_code)

    return jsonify({
        'success': success,
        'message': message
    })


@app.route('/api/settings/change_freigabe_code', methods=['POST'])
def api_change_freigabe_code():
    """API: Freigabe-Code ändern"""
    from managers.settings_manager import set_freigabe_code, verify_freigabe_code
    data = request.json
    current_code = data.get('current_code', '')
    new_code = data.get('new_code', '')

    if not verify_freigabe_code(current_code):
        return jsonify({'success': False, 'message': 'Aktueller Freigabe-Code ist falsch'})

    success, message = set_freigabe_code(new_code)
    return jsonify({'success': success, 'message': message})


@app.route('/admin_exam_config')
def admin_exam_config():
    """Admin: Prüfungsmodus-Konfiguration"""
    exam_settings = get_exam_settings()
    stromkreise = get_all_stromkreise()
    return render_template('admin_exam_config.html',
                           exam_settings=exam_settings,
                           stromkreise=stromkreise)


@app.route('/api/settings/exam', methods=['POST'])
def api_save_exam_settings():
    """API: Prüfungs-Einstellungen speichern"""
    data = request.json
    error_count = int(data.get('exam_error_count', 3))
    duration_minutes = int(data.get('exam_duration_minutes', 20))
    allowed_stromkreise = data.get('exam_allowed_stromkreise', [])

    success, message = set_exam_settings(error_count, duration_minutes, allowed_stromkreise)
    return jsonify({'success': success, 'message': message})


@app.route('/api/exam/force_end', methods=['POST'])
def api_exam_force_end():
    """Admin-Override: Prüfung zwangsweise beenden (erfordert Admin-Code)"""
    global exam_active, exam_client_ip, exam_number_current, exam_start_time
    data = request.json or {}
    code = data.get('code', '')
    if not verify_admin_code(code):
        return jsonify({'success': False, 'error': 'Falscher Admin-Code'}), 403
    relay_controller.reset_all_relays()
    exam_active = False
    exam_client_ip = None
    exam_number_current = None
    exam_start_time = None
    print(f"🔓 Prüfung per Admin-Override beendet von {get_client_ip()}")
    return jsonify({'success': True, 'message': 'Prüfung beendet, alle Relais zurückgesetzt.'})


@app.route('/shutdown_system', methods=['POST'])
def shutdown_system():
    """Fährt das System herunter"""
    try:
        relay_controller.reset_all_relays()
        subprocess.Popen(['sudo', 'shutdown', '-h', 'now'])
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/restart_system', methods=['POST'])
def restart_system():
    """Startet das System neu"""
    try:
        relay_controller.reset_all_relays()
        subprocess.Popen(['sudo', 'reboot'])
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


# ==================== DOKUMENTE / PDF-VIEWER ====================

PDF_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'pdfs')

@app.route('/api/pdfs', methods=['GET'])
def api_list_pdfs():
    """Listet alle PDFs und Excel-Dateien im pdfs/-Ordner"""
    import os as _os
    try:
        if not _os.path.exists(PDF_DIR):
            _os.makedirs(PDF_DIR)
        files = sorted([
            f for f in _os.listdir(PDF_DIR)
            if f.lower().endswith('.pdf') or f.lower().endswith('.xlsx')
        ])
        return jsonify({'success': True, 'files': files})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})


@app.route('/pdfs/<path:filename>')
def serve_pdf(filename):
    """Stellt eine PDF-Datei bereit"""
    import os as _os
    from flask import send_from_directory, abort
    # Sicherheit: nur .pdf und kein Pfad-Traversal
    if '..' in filename or '/' in filename or not filename.lower().endswith('.pdf'):
        abort(400)
    if not _os.path.exists(_os.path.join(PDF_DIR, filename)):
        abort(404)
    return send_from_directory(PDF_DIR, filename, mimetype='application/pdf')


@app.route('/documents/<path:filename>')
def serve_document(filename):
    """Stellt eine Datei aus dem pdfs/-Ordner zum Download bereit (PDF oder Excel)"""
    import os as _os
    from flask import send_from_directory, abort
    allowed = ('.pdf', '.xlsx')
    if '..' in filename or '/' in filename or not filename.lower().endswith(allowed):
        abort(400)
    if not _os.path.exists(_os.path.join(PDF_DIR, filename)):
        abort(404)
    return send_from_directory(PDF_DIR, filename, as_attachment=True)


# ==================== DEBUG/STATUS ====================

@app.route('/debug_status')
def debug_status():
    """Debug-Informationen"""
    from hardware.serial_handler import SERIAL_AVAILABLE
    
    return jsonify({
        'serial_available': SERIAL_AVAILABLE,
        'active_relays': relay_controller.active_relays,
        'relay_states_module_0': relay_controller.relay_states[0][:10],
        'relay_states_module_1': relay_controller.relay_states[1][:10],
        'stromkreise': {k: v['name'] for k, v in STROMKREISE.items()},
        'modbus_modules': MODBUS_MODULES
    })


# ==================== APP START ====================

def initialize_app(skip_gpio_check=False):
    """Initialisiert die App einmalig"""
    init_db()

    # Lade dynamische Gruppen und Namen beim Start
    import config
    config.RELAY_GROUPS = get_all_groups()
    config.RELAY_NAMES = get_all_relay_names()

    # Initialisiere GPIO-Monitor
    import os
    gpio_pin1 = getattr(config, 'GPIO_MONITOR_PIN1', 17)
    gpio_pin2 = getattr(config, 'GPIO_MONITOR_PIN2', 27)
    # Gespeicherten Timeout bevorzugen, config.py als Fallback
    from managers.settings_manager import get_gpio_shutdown_timeout
    gpio_shutdown_timeout = get_gpio_shutdown_timeout()

    # Bei Flask dev server: nur im Reloader
    # Bei Gunicorn: skip_gpio_check=True (wird vom Hook aufgerufen)
    if skip_gpio_check or os.environ.get('WERKZEUG_RUN_MAIN') == 'true' or not DEBUG:
        monitor = init_gpio_monitor(pin1=gpio_pin1, pin2=gpio_pin2, shutdown_timeout=gpio_shutdown_timeout)
        # Sicherheitsrelais beim Notaus automatisch abschalten
        def _notaus_callback():
            from config import SAFETY_RELAY_ID
            if relay_controller.safety_relay_state:
                print("🔴 Notaus: Sicherheitsrelais wird abgeschaltet")
                relay_controller.safety_relay_state = False
                relay_controller.set_relay(SAFETY_RELAY_ID, False)
        if monitor:
            monitor.on_notaus_active = _notaus_callback
    else:
        print("⏭️ GPIO-Monitor wird im Reloader-Prozess übersprungen")

    from hardware.serial_handler import SERIAL_AVAILABLE

    print("=" * 60)
    print("VDE Messwand - Modulare Version")
    print("=" * 60)
    print(f"Serial Port: {SERIAL_PORT}")
    print(f"Baud Rate: {BAUD_RATE}")
    print(f"Serial Status: {'✅ Hardware Ready' if SERIAL_AVAILABLE else '🔧 Dummy Mode'}")
    print(f"\nGPIO-Monitor: Pin {gpio_pin1} und {gpio_pin2}")
    print(f"Warnung: 'Notaus betätigt' bei geschlossenem Schließer")
    # Sicherheitsrelais beim Start sicher ausschalten
    from config import SAFETY_RELAY_ID
    relay_controller.safety_relay_state = False
    relay_controller.set_relay(SAFETY_RELAY_ID, False)
    print(f"\nRelais: 0-{SAFETY_RELAY_ID - 1} ({SAFETY_RELAY_ID} nutzbar, Relais {SAFETY_RELAY_ID + 1} = Sicherheitsrelais)")
    print(f"Relais-Gruppen: {len(config.RELAY_GROUPS)} definiert")
    print(f"Benannte Relais: {len(config.RELAY_NAMES)} definiert")
    print(f"Stromkreise für UI-Gruppierung:")
    for sk_num, sk_data in STROMKREISE.items():
        print(f"  {sk_num}. {sk_data['name']}")
    print("=" * 60)

    try:
        app.run(host=HOST, port=PORT, debug=DEBUG)
    finally:
        # Cleanup beim Beenden
        cleanup_gpio()


# ==================== GUNICORN HOOKS ====================

def on_starting(server):
    """Gunicorn Hook: Wird beim Start des Master-Prozesses aufgerufen (vor den Workern)"""
    print("🚀 Gunicorn Master-Prozess: Initialisiere App...")
    initialize_app(skip_gpio_check=True)
    print("✅ App-Initialisierung abgeschlossen")


if __name__ == '__main__':
    initialize_app()